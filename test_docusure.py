import io
import unittest
from datetime import date

import pandas as pd
import pymupdf as fitz
from openpyxl import load_workbook

from app import (
    FieldRecord,
    Finding,
    ProcessedDocument,
    apply_vessel_identity_gate,
    assign_document_groups,
    documents_for_group,
    extract_class_status,
    first_field,
    make_excel_v20,
    parse_date_any,
    process_document_bytes,
    valid_imo_number,
)
from docusure_engine import (
    DOC_LABELS,
    DocumentProfile,
    build_sts_screen,
    certificate_watch_df,
    detect_document_type,
    enhance_register,
    extract_mooring_summary,
)


def field(source, field_id, value, confidence_score=92):
    return FieldRecord(
        source=source,
        field_id=field_id,
        label=field_id,
        value=value,
        confidence="deterministic",
        confidence_score=confidence_score,
    )


def profile(doc_type, imo, filename=None, vessel_name="TEST VESSEL"):
    return DocumentProfile(
        filename=filename or f"{doc_type.lower()}.pdf",
        doc_type=doc_type,
        doc_label=DOC_LABELS[doc_type],
        confidence_score=95,
        confidence="High",
        classification_evidence="test fixture",
        imo=imo,
        vessel_name=vessel_name,
        pages=1,
        pages_with_text=1,
        text_chars=500,
        text_quality="Good — searchable text detected",
        assigned_group=imo,
    )


class DateAndIdentityTests(unittest.TestCase):
    def test_maritime_dates_are_day_first_and_question_numbers_are_not_dates(self):
        self.assertEqual(parse_date_any("05/08/2025"), date(2025, 8, 5))
        self.assertEqual(parse_date_any("05-Aug-2025"), date(2025, 8, 5))
        self.assertEqual(parse_date_any("2025-08-05"), date(2025, 8, 5))
        self.assertIsNone(parse_date_any("Question 10.1.7"))

    def test_imo_checksum_and_candidate_selection(self):
        self.assertTrue(valid_imo_number("IMO 1234567"))
        self.assertFalse(valid_imo_number("1234560"))
        fields = [
            field("HVPQ", "vessel.imo", "1234560", confidence_score=96),
            field("HVPQ", "vessel.imo", "1234567", confidence_score=35),
        ]
        self.assertEqual(first_field(fields, "HVPQ", "vessel.imo"), "1234567")

    def test_cross_vessel_comparisons_are_blocked(self):
        fields = [
            field("HVPQ", "vessel.imo", "9074729"),
            field("CLASS", "vessel.imo", "9303807"),
        ]
        findings = [Finding(
            area="Certificates",
            check="SMC expiry",
            status="MISMATCH",
            risk="HIGH",
            hvpq_value="2026-01-01",
            class_value="2027-01-01",
        )]
        gated, same_vessel, by_source = apply_vessel_identity_gate(findings, fields)
        self.assertFalse(same_vessel)
        self.assertEqual(set(by_source.values()), {"9074729", "9303807"})
        self.assertEqual(gated[0].status, "BLOCKED")


class ClassificationAndExtractionTests(unittest.TestCase):
    def test_document_classification_uses_content_signatures(self):
        hvpq_type, hvpq_score, _ = detect_document_type(
            "random-name.pdf",
            "Harmonized Vessel Particulars Questionnaire\nDate this HVPQ document completed 01-Sep-2026\nProvide details for mooring ropes",
        )
        sts_type, sts_score, _ = detect_document_type(
            "manual.pdf",
            "Ship to Ship Transfer Operations Plan\nMARPOL Annex I Regulation 41",
        )
        class_type, class_score, _ = detect_document_type(
            "report.pdf",
            "CLASS STATUS\nCertificate Description\nStatus of Surveys",
        )
        self.assertEqual((hvpq_type, sts_type, class_type), ("HVPQ", "STS_PLAN", "CLASS"))
        self.assertGreaterEqual(min(hvpq_score, sts_score, class_score), 80)

    def test_individual_certificate_pdf_is_classified_and_extracted(self):
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text(
            (72, 72),
            "SAFETY MANAGEMENT CERTIFICATE\n"
            "Name of Ship\nTEST VESSEL\n"
            "IMO Number: 9074729\n"
            "Certificate Number: SMC-12345\n"
            "Issued date: 01-Jan-2024\n"
            "Valid until: 01-Jan-2029\n"
            "This certificate remains subject to annual verification.",
        )
        data = pdf.tobytes()
        pdf.close()
        processed = process_document_bytes("vessel_smc_certificate.pdf", data, enable_ocr=False)
        self.assertEqual(processed.profile.doc_type, "CERTIFICATE")
        self.assertEqual(processed.profile.imo, "9074729")
        self.assertEqual(first_field(processed.fields, "CERTIFICATE", "cert.smc.expiry"), "2029-01-01")
        self.assertEqual(first_field(processed.fields, "CERTIFICATE", "cert.smc.number"), "SMC-12345")
        self.assertTrue(all(item.document_name == "vessel_smc_certificate.pdf" for item in processed.fields))

    def test_class_issue_date_is_not_mislabeled_as_expiry(self):
        pages = [(1, "CLASS STATUS\nSafety Management Certificate\nIssued date\n01-Jan-2024")]
        fields = extract_class_status(pages)
        self.assertEqual(first_field(fields, "CLASS", "cert.smc.issue"), "2024-01-01")
        self.assertEqual(first_field(fields, "CLASS", "cert.smc.expiry"), "")

    def test_class_headings_alone_do_not_create_open_items(self):
        fields = extract_class_status([(1, "CLASS STATUS\nConditions of Class\nMemoranda of Class\nDispensations")])
        self.assertEqual(first_field(fields, "CLASS", "classification.conditions_of_class"), "")
        self.assertEqual(first_field(fields, "CLASS", "classification.memo_of_class"), "")
        self.assertEqual(first_field(fields, "CLASS", "classification.flag_dispensation"), "")

    def test_class_nil_and_open_markers_are_distinguished(self):
        nil_fields = extract_class_status([(1, "CLASS STATUS\nConditions of Class\nNil")])
        open_fields = extract_class_status([(
            1,
            "CLASS STATUS\nConditions of Class\nCondition No. 12\nDue date 31-Dec-2026\nRepair shell insert",
        )])
        self.assertEqual(first_field(nil_fields, "CLASS", "classification.conditions_of_class"), "No")
        self.assertTrue(first_field(open_fields, "CLASS", "classification.conditions_of_class").startswith("Yes"))

    def test_mooring_heading_does_not_become_four_inventory_records(self):
        text = """
        10.1.3 Mooring design SDMBL 55 tonnes
        10.1.4 Mooring winch brake test Date of last brake test 01-Sep-2026
        10.1.7 Provide details for Mooring Ropes, Wires, Tails and Shackles
        Tail Aft Polyester 80 mm 11 m TDBF 70 tonnes Date installed 01-Jan-2026
        Rope Forward HMPE 72 mm 220 m LDBF 58 tonnes Date installed 02-Feb-2025
        10.1.8 Next section
        """
        summary, inventory = extract_mooring_summary({"HVPQ": text}, date(2026, 9, 18))
        self.assertEqual(inventory["Type"].tolist(), ["Tail", "Rope / line"])
        self.assertNotIn("Wire", inventory["Type"].tolist())
        self.assertNotIn("Shackle", inventory["Type"].tolist())
        values = dict(zip(summary["Item"], summary["Value"]))
        self.assertEqual(values["Latest visible mooring brake test"], "2026-09-01")
        self.assertEqual(values["SDMBL"], "55 tonnes")

    def test_uncertain_files_are_not_silently_mixed_with_a_different_named_vessel(self):
        alpha = ProcessedDocument(profile("HVPQ", "9074729", "alpha.pdf", "ALPHA"), [], "", [], "a")
        generic = ProcessedDocument(profile("CERTIFICATE", "", "smc.pdf", "SMC"), [], "", [], "b")
        bravo = ProcessedDocument(profile("CLASS", "", "bravo-class.pdf", "BRAVO"), [], "", [], "c")
        documents = [alpha, generic, bravo]
        assign_document_groups(documents)
        self.assertEqual(generic.profile.assigned_group, "9074729")
        self.assertEqual(generic.profile.grouping_confidence, "Low")
        self.assertEqual(bravo.profile.assigned_group, "Unassigned")

    def test_labelled_name_can_form_a_non_imo_single_vessel_group(self):
        named = ProcessedDocument(
            profile("HVPQ", "", "named-hvpq.pdf", "BRAVO"),
            [],
            "Name of Ship\nBRAVO\nHarmonized Vessel Particulars Questionnaire",
            [],
            "named",
        )
        assign_document_groups([named])
        self.assertEqual(named.profile.assigned_group, "Name: BRAVO")
        self.assertEqual(named.profile.grouping_confidence, "Medium")

    def test_joint_sts_assessment_is_available_to_both_imo_groups(self):
        joint_profile = profile("STS_ASSESSMENT", "9074729", "joint-jpo.pdf", "ALPHA")
        joint_profile.related_imos = "9074729, 9303807"
        joint = ProcessedDocument(joint_profile, [], "", [], "joint")
        alpha = ProcessedDocument(profile("HVPQ", "9074729", "alpha.pdf", "ALPHA"), [], "", [], "a")
        bravo = ProcessedDocument(profile("HVPQ", "9303807", "bravo.pdf", "BRAVO"), [], "", [], "b")
        documents = [joint, alpha, bravo]
        assign_document_groups(documents)
        self.assertIn(joint, documents_for_group(documents, "9074729"))
        self.assertIn(joint, documents_for_group(documents, "9303807"))
        self.assertIn("Shared STS evidence", joint.profile.grouping_evidence)


class DecisionOutputTests(unittest.TestCase):
    def test_certificate_watch_separates_expired_short_term_and_valid(self):
        fields = [
            field("HVPQ", "cert.smc.expiry", "01-Jan-2026", 80),
            field("CLASS", "cert.smc.expiry", "01-Oct-2026", 92),
            field("CERTIFICATE", "cert.issc.expiry", "01-Aug-2026", 96),
            field("CERTIFICATE", "cert.iopp.expiry", "01-Jan-2028", 96),
        ]
        watch = certificate_watch_df(fields, date(2026, 9, 18), horizon_days=180)
        by_cert = watch.set_index("Certificate")
        self.assertEqual(by_cert.loc["Safety Management Certificate", "Status"], "Expires within 30 days")
        self.assertEqual(by_cert.loc["Safety Management Certificate", "Preferred source"], "CLASS")
        self.assertEqual(by_cert.loc["International Ship Security Certificate", "Status"], "Expired")
        self.assertEqual(by_cert.loc["International Oil Pollution Prevention", "Status"], "Valid beyond watch window")

    def test_extraction_gap_is_not_reported_as_a_confirmed_error(self):
        source = pd.DataFrame([{
            "Priority": "HIGH",
            "Status": "Blank/not extracted",
            "Check": "Certificate expiry",
            "Action": "Check source",
        }])
        enhanced = enhance_register(source)
        self.assertEqual(enhanced.iloc[0]["Conclusion"], "Not verified")
        self.assertEqual(enhanced.iloc[0]["Evidence confidence"], "Low")

    def test_sts_open_condition_produces_hold_not_clearance(self):
        primary, counterpart = "9074729", "9303807"
        profiles = [
            profile("HVPQ", primary, "a_hvpq.pdf", "ALPHA"),
            profile("CLASS", primary, "a_class.pdf", "ALPHA"),
            profile("HVPQ", counterpart, "b_hvpq.pdf", "BRAVO"),
        ]
        fields = {
            primary: [field("CLASS", "classification.conditions_of_class", "Yes — Condition No. 12 open")],
            counterpart: [],
        }
        texts = {
            primary: ["LOA 250 m Beam 44 m cargo manifold 16 inch"],
            counterpart: ["LOA 230 m Beam 42 m cargo manifold 16 inch"],
        }
        decision, screen, facts, _ = build_sts_screen(
            profiles, fields, texts, primary, counterpart, date(2026, 9, 18)
        )
        self.assertEqual(decision, "HOLD")
        self.assertTrue(((screen["Status"] == "HOLD") & (screen["Check"] == "Conditions of Class")).any())
        self.assertIn("Common nominal size visible", facts.loc[facts["Compatibility fact"] == "Visible manifold sizes", "Screening result"].iloc[0])

    def test_excel_register_contains_styled_named_sheets(self):
        data = make_excel_v20([
            ("Decision", pd.DataFrame([{"Status": "HOLD", "Reason": "Expired"}])),
            ("Certificate Watch", pd.DataFrame([{"Status": "Valid", "Days remaining": 365}])),
        ])
        workbook = load_workbook(io.BytesIO(data))
        self.assertEqual(workbook.sheetnames, ["Decision", "Certificate Watch"])
        self.assertEqual(workbook["Decision"].freeze_panes, "A2")
        self.assertEqual(workbook["Decision"]["A1"].fill.fgColor.rgb[-6:], "17365D")


if __name__ == "__main__":
    unittest.main()
